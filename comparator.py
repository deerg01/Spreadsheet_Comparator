import time
import threading
import csv
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

from openpyxl import load_workbook

from file_handler import find_pairs
from excel_handler import (
    copy_new_file_to_output,
    create_csv_output_as_xlsx,
    apply_highlights,
    open_output_folder,
)


def normalize_value(value):
    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    return str(value).strip()


def is_empty(value):
    return normalize_value(value) == ""


def send_progress(progress_callback, message):
    if progress_callback is not None:
        progress_callback(message)


def format_elapsed_time(seconds):
    seconds = int(seconds)

    hours, remainder = divmod(seconds, 3600)
    minutes, seconds = divmod(remainder, 60)

    if hours:
        return f"{hours}h {minutes}m {seconds}s"

    if minutes:
        return f"{minutes}m {seconds}s"

    return f"{seconds}s"


def get_xlsx_rows(ws):
    rows = []

    for row_number in range(1, ws.max_row + 1):
        row = []

        for column_number in range(1, ws.max_column + 1):
            row.append(
                ws.cell(
                    row=row_number,
                    column=column_number,
                ).value
            )

        rows.append(row)

    return rows


def read_csv_rows(file_path):
    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        return list(csv.reader(file))


def row_has_value(row):
    if row is None:
        return False

    for value in row:
        if not is_empty(value):
            return True

    return False


def rows_are_both_empty(old_row, new_row):
    old_empty = not row_has_value(old_row)
    new_empty = not row_has_value(new_row)

    return old_empty and new_empty


def rows_match(old_row, new_row):
    if old_row is None or new_row is None:
        return False

    if not row_has_value(old_row):
        return not row_has_value(new_row)

    max_columns = max(
        len(old_row),
        len(new_row),
    )

    for column in range(max_columns):
        old_value = (
            old_row[column]
            if column < len(old_row)
            else None
        )

        new_value = (
            new_row[column]
            if column < len(new_row)
            else None
        )

        if is_empty(old_value):
            continue

        if is_empty(new_value):
            return False

        if normalize_value(old_value) != normalize_value(new_value):
            return False

    return True


def find_future_matching_row(
    source_rows,
    target_rows,
    source_index,
    target_start_index,
    max_distance=20,
):
    source_row = source_rows[source_index]

    if not row_has_value(source_row):
        return None

    end_index = min(
        len(target_rows),
        target_start_index + max_distance + 1,
    )

    for target_index in range(
        target_start_index,
        end_index,
    ):
        target_row = target_rows[target_index]

        if rows_match(
            source_row,
            target_row,
        ):
            return target_index

    return None


def compare_mapped_row(
    old_row,
    new_row,
    new_row_number,
):
    differences = []

    max_columns = max(
        len(old_row),
        len(new_row),
    )

    for column in range(max_columns):
        old_value = (
            old_row[column]
            if column < len(old_row)
            else None
        )

        new_value = (
            new_row[column]
            if column < len(new_row)
            else None
        )

        if is_empty(old_value):
            continue

        if is_empty(new_value):
            differences.append(
                (
                    new_row_number,
                    column + 1,
                )
            )
            continue

        if normalize_value(old_value) != normalize_value(new_value):
            differences.append(
                (
                    new_row_number,
                    column + 1,
                )
            )

    return differences


def compare_rows(
    old_rows,
    new_rows,
    max_lookahead=20,
):
    differences = []
    added_rows = []
    missing_rows = []

    old_index = 0
    new_index = 0

    while (
        old_index < len(old_rows)
        and new_index < len(new_rows)
    ):
        old_row = old_rows[old_index]
        new_row = new_rows[new_index]

        if rows_are_both_empty(
            old_row,
            new_row,
        ):
            old_index += 1
            new_index += 1
            continue

        if rows_match(
            old_row,
            new_row,
        ):
            old_index += 1
            new_index += 1
            continue

        future_new_index = find_future_matching_row(
            old_rows,
            new_rows,
            old_index,
            new_index + 1,
            max_lookahead,
        )

        future_old_index = find_future_matching_row(
            new_rows,
            old_rows,
            new_index,
            old_index + 1,
            max_lookahead,
        )

        if (
            future_new_index is not None
            and future_old_index is None
        ):
            for index in range(
                new_index,
                future_new_index,
            ):
                if row_has_value(
                    new_rows[index]
                ):
                    added_rows.append(index)

            new_index = future_new_index
            continue

        if (
            future_old_index is not None
            and future_new_index is None
        ):
            for index in range(
                old_index,
                future_old_index,
            ):
                if row_has_value(
                    old_rows[index]
                ):
                    missing_rows.append(index)

            old_index = future_old_index
            continue

        if (
            future_new_index is not None
            and future_old_index is not None
        ):
            new_distance = (
                future_new_index - new_index
            )

            old_distance = (
                future_old_index - old_index
            )

            if new_distance < old_distance:
                for index in range(
                    new_index,
                    future_new_index,
                ):
                    if row_has_value(
                        new_rows[index]
                    ):
                        added_rows.append(index)

                new_index = future_new_index
                continue

            if old_distance < new_distance:
                for index in range(
                    old_index,
                    future_old_index,
                ):
                    if row_has_value(
                        old_rows[index]
                    ):
                        missing_rows.append(index)

                old_index = future_old_index
                continue

            differences.extend(
                compare_mapped_row(
                    old_row,
                    new_row,
                    new_index + 1,
                )
            )

            old_index += 1
            new_index += 1
            continue

        differences.extend(
            compare_mapped_row(
                old_row,
                new_row,
                new_index + 1,
            )
        )

        old_index += 1
        new_index += 1

    while new_index < len(new_rows):
        if row_has_value(
            new_rows[new_index]
        ):
            added_rows.append(new_index)

        new_index += 1

    while old_index < len(old_rows):
        if row_has_value(
            old_rows[old_index]
        ):
            missing_rows.append(old_index)

        old_index += 1

    return (
        differences,
        added_rows,
        missing_rows,
    )


def compare_xlsx_files(
    old_file,
    new_file,
    cancel_event=None,
):
    old_workbook = load_workbook(
        old_file,
        read_only=False,
        data_only=False,
    )

    new_workbook = load_workbook(
        new_file,
        read_only=False,
        data_only=False,
    )

    differences = {}
    added_rows = {}
    missing_rows = {}

    for sheet_name in new_workbook.sheetnames:
        if (
            cancel_event is not None
            and cancel_event.is_set()
        ):
            raise InterruptedError

        new_ws = new_workbook[sheet_name]

        if sheet_name in old_workbook.sheetnames:
            old_ws = old_workbook[sheet_name]

            old_rows = get_xlsx_rows(
                old_ws
            )

            new_rows = get_xlsx_rows(
                new_ws
            )

            (
                sheet_differences,
                sheet_added_rows,
                sheet_missing_rows,
            ) = compare_rows(
                old_rows,
                new_rows,
            )

            differences[sheet_name] = (
                sheet_differences
            )

            added_rows[sheet_name] = (
                sheet_added_rows
            )

            missing_rows[sheet_name] = (
                sheet_missing_rows
            )

        else:
            sheet_added_rows = []

            for index, row in enumerate(
                get_xlsx_rows(new_ws)
            ):
                if row_has_value(row):
                    sheet_added_rows.append(index)

            differences[sheet_name] = []
            added_rows[sheet_name] = (
                sheet_added_rows
            )
            missing_rows[sheet_name] = []

    for sheet_name in old_workbook.sheetnames:
        if sheet_name in new_workbook.sheetnames:
            continue

        old_ws = old_workbook[sheet_name]

        old_rows = get_xlsx_rows(
            old_ws
        )

        missing_rows[sheet_name] = [
            index
            for index, row in enumerate(old_rows)
            if row_has_value(row)
        ]

        differences[sheet_name] = []
        added_rows[sheet_name] = []

    return (
        differences,
        added_rows,
        missing_rows,
        old_workbook,
    )


def compare_csv_files(
    old_file,
    new_file,
    cancel_event=None,
):
    if (
        cancel_event is not None
        and cancel_event.is_set()
    ):
        raise InterruptedError

    old_rows = read_csv_rows(
        old_file
    )

    new_rows = read_csv_rows(
        new_file
    )

    (
        differences,
        added_rows,
        missing_rows,
    ) = compare_rows(
        old_rows,
        new_rows,
    )

    return (
        {
            "Sheet1": differences
        },
        {
            "Sheet1": added_rows
        },
        {
            "Sheet1": missing_rows
        },
        old_rows,
    )


def count_differences(differences):
    return sum(
        len(cells)
        for cells in differences.values()
    )


def process_pair(
    pair,
    output_dir,
    current_index,
    total_count,
    cancel_event=None,
    progress_callback=None,
):
    if (
        cancel_event is not None
        and cancel_event.is_set()
    ):
        raise InterruptedError

    old_file, new_file = pair

    old_file = Path(old_file)
    new_file = Path(new_file)

    output_dir = Path(output_dir)

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    output_file = (
        output_dir
        / f"{new_file.stem}_comparison.xlsx"
    )

    send_progress(
        progress_callback,
        f"[{current_index:02d}/{total_count}] Comparing: {new_file.name}",
    )

    if new_file.suffix.lower() == ".xlsx":
        copy_new_file_to_output(
            new_file,
            output_file,
        )

        (
            differences,
            added_rows,
            missing_rows,
            old_source,
        ) = compare_xlsx_files(
            old_file,
            new_file,
            cancel_event,
        )

    elif new_file.suffix.lower() == ".csv":
        create_csv_output_as_xlsx(
            new_file,
            output_file,
        )

        (
            differences,
            added_rows,
            missing_rows,
            old_source,
        ) = compare_csv_files(
            old_file,
            new_file,
            cancel_event,
        )

    else:
        raise ValueError(
            f"Unsupported file type: {new_file.suffix}"
        )

    if (
        cancel_event is not None
        and cancel_event.is_set()
    ):
        raise InterruptedError

    apply_highlights(
        output_file,
        differences,
        added_rows,
        missing_rows,
        old_source,
    )

    difference_count = count_differences(
        differences
    )

    added_count = sum(
        len(rows)
        for rows in added_rows.values()
    )

    missing_count = sum(
        len(rows)
        for rows in missing_rows.values()
    )

    send_progress(
        progress_callback,
        f"[{current_index:02d}/{total_count}] Completed: "
        f"{new_file.name} "
        f"({difference_count} cell differences, "
        f"{added_count} added rows, "
        f"{missing_count} missing rows)",
    )

    return {
        "old_file": old_file,
        "new_file": new_file,
        "status": "completed",
        "differences": differences,
        "difference_count": difference_count,
        "cell_differences": difference_count,
        "structural_differences": (
            added_count + missing_count
        ),
        "added_rows": added_rows,
        "missing_rows": missing_rows,
        "output_file": output_file,
        "output_path": output_file,
    }


def compare_versions(
    old_folder,
    new_folder,
    pairing_method="filename",
    parallel=False,
    progress_callback=None,
    cancel_event=None,
):
    start_time = time.time()

    old_folder = Path(old_folder)
    new_folder = Path(new_folder)

    if cancel_event is None:
        cancel_event = threading.Event()

    send_progress(
        progress_callback,
        "Scanning Old Version folder...",
    )

    send_progress(
        progress_callback,
        "Scanning New Version folder...",
    )

    pairs, unmatched_old, unmatched_new = find_pairs(
        old_folder,
        new_folder,
        pairing_method,
    )

    output_dir = (
        new_folder
        / "comparison_output"
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    total_count = len(pairs)

    if total_count == 0:
        send_progress(
            progress_callback,
            "[00/00] No matching files found.",
        )

        send_progress(
            progress_callback,
            f"Total elapsed time: "
            f"{format_elapsed_time(time.time() - start_time)}",
        )

        open_output_folder(
            output_dir
        )

        return {
            "results": [],
            "output_dir": output_dir,
            "unmatched_old": unmatched_old,
            "unmatched_new": unmatched_new,
        }

    results = []

    if parallel:
        max_workers = min(
            4,
            total_count,
        )

        with ThreadPoolExecutor(
            max_workers=max_workers
        ) as executor:
            future_map = {}

            for index, pair in enumerate(
                pairs,
                start=1,
            ):
                if cancel_event.is_set():
                    raise InterruptedError

                future = executor.submit(
                    process_pair,
                    pair,
                    output_dir,
                    index,
                    total_count,
                    cancel_event,
                    progress_callback,
                )

                future_map[future] = pair

            for future in as_completed(
                future_map
            ):
                if cancel_event.is_set():
                    for pending in future_map:
                        pending.cancel()

                    raise InterruptedError

                try:
                    result = future.result()
                    results.append(result)

                except InterruptedError:
                    cancel_event.set()

                    for pending in future_map:
                        pending.cancel()

                    raise

                except Exception as error:
                    old_file, new_file = (
                        future_map[future]
                    )

                    results.append(
                        {
                            "old_file": old_file,
                            "new_file": new_file,
                            "status": "error",
                            "differences": {},
                            "difference_count": 0,
                            "cell_differences": 0,
                            "structural_differences": 0,
                            "error": str(error),
                            "output_file": None,
                            "output_path": None,
                        }
                    )

                    send_progress(
                        progress_callback,
                        f"ERROR: {new_file.name} - {error}",
                    )

    else:
        for index, pair in enumerate(
            pairs,
            start=1,
        ):
            if cancel_event.is_set():
                raise InterruptedError

            try:
                result = process_pair(
                    pair,
                    output_dir,
                    index,
                    total_count,
                    cancel_event,
                    progress_callback,
                )

                results.append(result)

            except InterruptedError:
                raise

            except Exception as error:
                old_file, new_file = pair

                results.append(
                    {
                        "old_file": old_file,
                        "new_file": new_file,
                        "status": "error",
                        "differences": {},
                        "difference_count": 0,
                        "cell_differences": 0,
                        "structural_differences": 0,
                        "error": str(error),
                        "output_file": None,
                        "output_path": None,
                    }
                )

                send_progress(
                    progress_callback,
                    f"ERROR: {new_file.name} - {error}",
                )

    results.sort(
        key=lambda item: str(
            item.get(
                "new_file",
                "",
            )
        ).lower()
    )

    completed_count = sum(
        1
        for result in results
        if result.get("status")
        == "completed"
    )

    error_count = sum(
        1
        for result in results
        if result.get("status")
        == "error"
    )

    total_differences = sum(
        result.get(
            "difference_count",
            0,
        )
        for result in results
    )

    total_structural_differences = sum(
        result.get(
            "structural_differences",
            0,
        )
        for result in results
    )

    elapsed = time.time() - start_time

    send_progress(
        progress_callback,
        f"Finished: "
        f"{completed_count}/{total_count} file(s), "
        f"{total_differences} cell differences, "
        f"{total_structural_differences} structural differences, "
        f"{error_count} error(s)",
    )

    send_progress(
        progress_callback,
        f"Total elapsed time: "
        f"{format_elapsed_time(elapsed)}",
    )

    if not cancel_event.is_set():
        open_output_folder(
            output_dir
        )

    return {
        "results": results,
        "output_dir": output_dir,
        "unmatched_old": unmatched_old,
        "unmatched_new": unmatched_new,
    }