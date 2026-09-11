from pathlib import Path
import csv
import os
import shutil
import subprocess
import sys
from copy import copy

import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


RED = "F4CCCC"
BLUE = "9FC5E8"


def copy_new_file_to_output(source_file, output_file):
    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_file, output_file)


def create_csv_output_as_xlsx(source_file, output_file):
    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    with open(
        source_file,
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:
        reader = csv.reader(f)

        for row_number, row in enumerate(reader, start=1):
            for column_number, value in enumerate(row, start=1):
                ws.cell(
                    row=row_number,
                    column=column_number,
                    value=value,
                )

    wb.save(output_file)


def apply_fill(cell, color):
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor=color,
    )


def highlight_row(ws, row_number, color):
    max_column = max(ws.max_column, 1)

    for column in range(1, max_column + 1):
        apply_fill(
            ws.cell(
                row=row_number,
                column=column,
            ),
            color,
        )


def copy_cell_style(source_cell, target_cell):
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)

    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format

    if source_cell.alignment:
        target_cell.alignment = copy(source_cell.alignment)

    if source_cell.font:
        target_cell.font = copy(source_cell.font)

    if source_cell.border:
        target_cell.border = copy(source_cell.border)

    if source_cell.protection:
        target_cell.protection = copy(source_cell.protection)


def append_missing_xlsx_rows(
    ws,
    old_ws,
    missing_rows,
):
    if not missing_rows:
        return

    start_row = ws.max_row + 1

    for offset, old_row_index in enumerate(missing_rows):
        output_row_number = start_row + offset
        source_row_number = old_row_index + 1

        for column in range(1, old_ws.max_column + 1):
            source_cell = old_ws.cell(
                row=source_row_number,
                column=column,
            )

            target_cell = ws.cell(
                row=output_row_number,
                column=column,
                value=source_cell.value,
            )

            copy_cell_style(
                source_cell,
                target_cell,
            )

        highlight_row(
            ws,
            output_row_number,
            RED,
        )


def append_missing_csv_rows(
    ws,
    old_rows,
    missing_rows,
):
    if not missing_rows:
        return

    start_row = ws.max_row + 1

    for offset, old_row_index in enumerate(missing_rows):
        output_row_number = start_row + offset
        old_row = old_rows[old_row_index]

        for column, value in enumerate(old_row, start=1):
            ws.cell(
                row=output_row_number,
                column=column,
                value=value,
            )

        highlight_row(
            ws,
            output_row_number,
            RED,
        )


def highlight_added_rows(ws, added_rows):
    for row_index in added_rows:
        row_number = row_index + 1

        if row_number <= ws.max_row:
            highlight_row(
                ws,
                row_number,
                BLUE,
            )


def highlight_differences(ws, differences):
    for row_number, column_number in differences:
        if row_number <= ws.max_row:
            apply_fill(
                ws.cell(
                    row=row_number,
                    column=column_number,
                ),
                RED,
            )


def apply_highlights(
    output_file,
    differences,
    added_rows,
    missing_rows,
    old_source,
):
    output_file = Path(output_file)

    wb = load_workbook(
        output_file,
        data_only=False,
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        sheet_differences = differences.get(
            sheet_name,
            [],
        )

        sheet_added_rows = added_rows.get(
            sheet_name,
            [],
        )

        sheet_missing_rows = missing_rows.get(
            sheet_name,
            [],
        )

        highlight_differences(
            ws,
            sheet_differences,
        )

        highlight_added_rows(
            ws,
            sheet_added_rows,
        )

        if isinstance(old_source, openpyxl.Workbook):
            if sheet_name in old_source.sheetnames:
                old_ws = old_source[sheet_name]

                append_missing_xlsx_rows(
                    ws,
                    old_ws,
                    sheet_missing_rows,
                )

        elif isinstance(old_source, list):
            append_missing_csv_rows(
                ws,
                old_source,
                sheet_missing_rows,
            )

    wb.save(output_file)

    if isinstance(old_source, openpyxl.Workbook):
        old_source.close()


def write_comparison(
    output_file,
    differences,
):
    wb = load_workbook(
        output_file,
        data_only=False,
    )

    for sheet_name, sheet_differences in differences.items():
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]

        for row_number, column_number in sheet_differences:
            apply_fill(
                ws.cell(
                    row=row_number,
                    column=column_number,
                ),
                RED,
            )

    wb.save(output_file)


def open_output_folder(folder):
    folder = Path(folder)

    if not folder.exists():
        return

    if sys.platform.startswith("win"):
        os.startfile(folder)
    elif sys.platform == "darwin":
        subprocess.Popen(
            ["open", str(folder)]
        )
    else:
        subprocess.Popen(
            ["xdg-open", str(folder)]
        )